from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_guide_doc():
    doc = Document()

    # 한글 폰트 설정 (맑은 고딕)
    style = doc.styles['Normal']
    style.font.name = 'Malgun Gothic'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

    # 1. 제목
    title = doc.add_heading('AI 협업 아키텍트 가이드', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('부제: 코딩을 몰라도 프로그램을 만드는 비전공자를 위한 방법론')
    doc.add_paragraph('-' * 50)

    # 2. 마인드셋
    doc.add_heading('1. 마인드셋의 전환: 코더에서 아키텍트로', level=1)
    p = doc.add_paragraph()
    p.add_run('핵심: ').bold = True
    p.add_run('직접 벽돌을 쌓는 "시공자(Coder)"가 되려 하지 말고, AI에게 도면을 지시하는 "설계자(Architect)"가 되십시오.')

    doc.add_paragraph('• 기존 방식: 문법 암기 -> 타자 치기 -> 에러와 싸우기')
    doc.add_paragraph('• AI 협업 방식: 구조 구상 -> AI에게 지시 -> 코드 리뷰 및 조립')

    # 3. 3단계 전략 표
    doc.add_heading('2. 현실적인 개발 3단계 전략 (진화 과정)', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'

    hdr = table.rows[0].cells
    hdr[0].text = '단계'
    hdr[1].text = '별명 (목표)'
    hdr[2].text = '특징 및 행동 요령'

    data = [
        ('1단계', '천막 짓기\n(Make it Work)', '• 목표: 일단 돌아가게 만들기\n• 특징: 파일 하나에 다 넣기 (막 코딩)\n• AI에게: "일단 작동하게 짜줘"'),
        ('2단계', '판잣집 짓기\n(Make it Right)', '• 목표: 가독성 챙기기\n• 특징: 함수(def)로 기능 쪼개기\n• AI에게: "함수별로 정리해줘"'),
        ('3단계', '벽돌집 짓기\n(Make it Fast/OOP)', '• 목표: 유지보수와 확장성\n• 특징: 클래스(Class)로 구조화\n• AI에게: "클래스 구조로 바꿔줘"')
    ]

    for step, name, desc in data:
        row = table.add_row().cells
        row[0].text = step
        row[1].text = name
        row[2].text = desc

    doc.add_paragraph('')

    # 4. 학습 루틴
    doc.add_heading('3. 실전 학습 루틴: 샌드위치 기법', level=1)
    doc.add_paragraph('AI에게 100% 맡기지 말고, [기획 - AI 구현 - 인간 검토]의 과정을 거치세요.')

    doc.add_paragraph('1. 기획: "이미지 변환기를 만들자" (아이디어 구상)', style='List Number')
    doc.add_paragraph('2. 초안: "파이썬으로 짜줘" (1단계 코드 확보)', style='List Number')
    doc.add_paragraph('3. 구조화: "클래스 구조로 리팩토링해줘" (3단계 코드 확보)', style='List Number')
    doc.add_paragraph('4. 응용: "버튼 색깔을 바꿔볼까?" (직접 코드 위치 찾아 수정하기)', style='List Number')

    # 5. 결론
    doc.add_heading('4. 결론', level=1)
    doc.add_paragraph('당신은 부족한 개발자가 아닙니다. 전체 구조를 이해하고 결과물을 만들어내는 "스마트한 기획자"입니다.')
    doc.add_paragraph('이 방법론을 통해 더 많은 프로그램을 설계하고 만들어보세요.')

    # 저장
    filename = 'AI_Architect_Guide.docx'
    doc.save(filename)
    print(f"문서가 생성되었습니다: {filename}")

if __name__ == "__main__":
    create_guide_doc()
