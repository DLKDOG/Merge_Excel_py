"""
Refactored version (Effective Python 관점):
- Tk root 1개만 사용
- Worker thread에서 Tkinter UI 직접 접근 금지 (queue + root.after로 UI 업데이트)
- logging 설정 1회만
- 엑셀 병합 로직(Core) / IO / UI 분리 형태
- concat 반복 최적화 (list에 모아서 1회 concat)
- ExcelFile 객체 재사용
- 예외는 경계에서 잡고, 내부는 raise + logging.exception 중심
"""

from __future__ import annotations

import os
import sys
import time
import threading
import queue
import logging
from dataclasses import dataclass
from typing import List, Tuple, Optional

import pandas as pd  # type: ignore
import numpy as np  # type: ignore

import tkinter as tk
from tkinter import filedialog, simpledialog, messagebox, ttk

from openpyxl import load_workbook  # type: ignore
from openpyxl.chart import ScatterChart, Reference  # type: ignore

import plotly.express as px  # type: ignore


# -------------------------
# Logging
# -------------------------

LOG_FILE = "merge_excel_log.txt"


def setup_logging() -> None:
    # basicConfig는 1회만 의미 있으므로, 프로그램 시작 시 딱 1번 호출
    logging.basicConfig(
        filename=LOG_FILE,
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )
    logging.info("Logging setup complete.")


logger = logging.getLogger(__name__)


# -------------------------
# Utilities / Config
# -------------------------

@dataclass(frozen=True)
class MergeResult:
    merged_file_path: str
    sorted_file_path: Optional[str]
    excluded_sheets: List[str]
    dynamic_chart_line: Optional[str]
    dynamic_chart_sorted: Optional[str]


def generate_simple_filename(base_name: str, index: int) -> str:
    return f"{base_name}_{index}.xlsx"


def convert_and_clean_date_column(df: pd.DataFrame) -> pd.DataFrame:
    """Date 컬럼이 있으면 datetime 변환 + 결측 보정."""
    if "Date" in df.columns:
        try:
            df = df.copy()
            df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
            df["Date"] = df["Date"].ffill().bfill()
        except Exception:
            logger.exception("Error processing Date column")
    return df


# -------------------------
# Core: merge logic
# -------------------------

def load_merge_frames_from_files(
    source_files: List[str],
    progress_cb=None,  # progress_cb(step_done:int, total_steps:int, message:str) -> None
) -> Tuple[pd.DataFrame, List[str]]:
    """
    여러 엑셀 파일의 여러 시트에서,
    - 'Line' 컬럼이 존재하고
    - 'Line'이 numeric dtype 인 경우만
    Line을 index로 설정하여 병합 프레임 리스트로 모은 뒤 1회 concat.
    """
    frames: List[pd.DataFrame] = []
    excluded: List[str] = []

    total_files = len(source_files)
    for idx, file_path in enumerate(source_files, start=1):
        logger.info("Processing file: %s", file_path)

        try:
            xls = pd.ExcelFile(file_path)
            for sheet_name in xls.sheet_names:
                try:
                    df = pd.read_excel(xls, sheet_name=sheet_name)

                    if "Line" in df.columns and pd.api.types.is_numeric_dtype(df["Line"]):
                        df2 = df.set_index("Line")
                        frames.append(df2)
                    else:
                        excluded.append(f"{sheet_name} in {os.path.basename(file_path)}")
                        logger.info(
                            "Excluded sheet: %s (missing numeric 'Line')",
                            f"{sheet_name} in {os.path.basename(file_path)}",
                        )
                except Exception:
                    logger.exception("Failed reading sheet %s from %s", sheet_name, file_path)

        except Exception:
            logger.exception("Failed opening Excel file: %s", file_path)

        if progress_cb:
            progress_cb(idx, total_files, f"파일 처리 중: {idx}/{total_files}")

    if not frames:
        raise ValueError("유효한 시트를 찾지 못했습니다. (numeric 'Line' 컬럼이 있는 시트가 없음)")

    all_data = pd.concat(frames, axis=0)
    # index(Line) 기준으로 정렬 후, 중복 Line은 마지막 값으로 처리
    all_data = all_data.sort_index().groupby(level=0).last().reset_index()

    # Date 정리(한 곳에서만)
    all_data = convert_and_clean_date_column(all_data)

    return all_data, excluded


def sort_and_reorder_by_column(df: pd.DataFrame, sort_criteria: str) -> pd.DataFrame:
    """
    sort_criteria로 정렬하고, 해당 컬럼을 첫 번째 컬럼으로 이동.
    sort 컬럼 결측은 ffill/bfill 처리.
    """
    if sort_criteria not in df.columns:
        raise KeyError(f"열 '{sort_criteria}'을(를) 찾을 수 없습니다.")

    df2 = df.copy()
    df2 = df2.sort_values(by=sort_criteria, na_position="last")
    df2[sort_criteria] = df2[sort_criteria].ffill().bfill()

    cols = df2.columns.tolist()
    cols.insert(0, cols.pop(cols.index(sort_criteria)))
    df2 = df2[cols]
    return df2


# -------------------------
# IO: Excel chart & Plotly
# -------------------------

def add_chart_to_excel(file_path: str) -> None:
    """정렬된 엑셀 파일에 openpyxl ScatterChart 시트를 추가."""
    logger.info("Loading workbook for chart: %s", file_path)

    wb = load_workbook(file_path)
    ws = wb.active

    if ws is None or ws.max_row <= 2 or ws.max_column <= 1:
        logger.warning("Worksheet seems empty or not suitable for chart: %s", file_path)
        return

    chart = ScatterChart()
    chart.title = "Scatter Plot"
    chart.style = 13
    chart.x_axis.title = "X"
    chart.y_axis.title = "Values"

    # X축: 첫 컬럼(1열), 2행부터
    x_values = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)

    # Y축: 2열~마지막열
    # openpyxl ScatterChart에서 titles_from_data=True는 add_data에 header가 있어야 제대로 동작
    # 현재 데이터는 헤더가 1행에 있으므로 add_data로 한 열씩 넣는 구조 유지
    for col in range(2, ws.max_column + 1):
        y_values = Reference(ws, min_col=col, min_row=1, max_row=ws.max_row)  # header 포함
        chart.add_data(y_values, titles_from_data=True)
    chart.set_categories(x_values)

    # Chart sheet 생성(중복 이름 처리)
    chart_sheet_name = "Chart"
    if chart_sheet_name in wb.sheetnames:
        # 기존 시트가 있으면 삭제 후 재생성
        del wb[chart_sheet_name]
    chart_sheet = wb.create_sheet(title=chart_sheet_name)
    chart_sheet.add_chart(chart, "A1")

    wb.save(file_path)
    logger.info("Chart added and workbook saved: %s", file_path)


def plot_from_dataframe(
    df: pd.DataFrame,
    x_axis_column: str,
    title: str,
    destination_folder: str,
    out_name: str,
) -> str:
    """
    plotly html 저장.
    기존 동작 호환:
    - 숫자형 컬럼들을 y로 여러 시리즈 생성
    - Date가 있으면 y2로 추가 (원래 코드 유지)
    """
    numeric_columns = df.select_dtypes(include="number").columns.tolist()

    # X축으로 사용하는 컬럼은 Y시리즈에서 제외
    if x_axis_column in numeric_columns:
        numeric_columns.remove(x_axis_column)

    # Date -> 오른쪽 y축(기존 동작 유지)
    if "Date" in df.columns:
        df = df.copy()
        df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
        right_y_column = "Date"
    else:
        right_y_column = None

    fig = px.scatter()

    for col in numeric_columns:
        fig.add_scatter(x=df[x_axis_column], y=df[col], mode="lines", name=col, yaxis="y1")

    if right_y_column and x_axis_column != "Date":
        fig.add_scatter(x=df[x_axis_column], y=df[right_y_column], mode="lines", name="Date", yaxis="y2")
        fig.update_layout(
            title=title,
            xaxis=dict(title=x_axis_column),
            yaxis=dict(title="Numeric Values", side="left"),
            yaxis2=dict(title="Date", overlaying="y", side="right"),
        )
    else:
        fig.update_layout(
            title=title,
            xaxis=dict(title=x_axis_column),
            yaxis=dict(title="Numeric Values"),
        )

    out_path = os.path.join(destination_folder, out_name)
    fig.write_html(out_path)
    try:
        fig.show()
    except Exception:
        # headless 환경 등에서 show 실패할 수 있으니 무시
        logger.exception("fig.show() failed (ignored)")
    return out_path


# -------------------------
# UI helpers (single root)
# -------------------------


def select_files(root: tk.Tk, title: str) -> List[str]:
    messagebox.showinfo(
        "파일 선택",
        "여러 파일을 선택하려면 Ctrl 또는 Shift 키를 누르고 선택하세요.",
        parent=root,
    )
    files_selected = filedialog.askopenfilenames(
        parent=root,
        title=title,
        filetypes=[("Excel files", "*.xlsx")],
    )
    # askopenfilenames can return an empty string if cancelled
    if isinstance(files_selected, str):
        return []
    return list(files_selected)


def select_folder(root: tk.Tk, title: str) -> str:
    folder = filedialog.askdirectory(parent=root, title=title)
    return folder


def select_sort_criteria(root: tk.Tk) -> str:
    sort_criteria = simpledialog.askstring(
        "정렬 기준",
        "정렬 기준을 입력하세요 (정렬 기준열 이름을 그대로 입력):",
        parent=root,
    )
    return sort_criteria or ""


# -------------------------
# Worker thread + UI queue
# -------------------------

@dataclass
class ProgressEvent:
    percent: int
    message: str


def merge_pipeline(
    source_files: List[str],
    destination_folder: str,
    sort_criteria: str,
    progress_q: "queue.Queue[ProgressEvent]",
) -> MergeResult:
    """
    워커 스레드에서 실행되는 파이프라인.
    UI는 만지지 않고 progress_q로만 이벤트 전달.
    """
    base_name = "merged_file"

    def push_progress(pct: int, msg: str) -> None:
        progress_q.put(ProgressEvent(percent=pct, message=msg))

    # Step model (대략적인 진행률)
    # 1) 파일 로드/병합: 0~70
    # 2) 저장(merged): 70~80
    # 3) 정렬/저장(sorted): 80~90
    # 4) openpyxl 차트 + plotly: 90~100
    total_files = max(len(source_files), 1)

    def progress_cb(done: int, total: int, message: str) -> None:
        pct = int((done / total) * 70)
        push_progress(pct, message)

    # 1) Merge
    push_progress(0, "병합 준비 중...")
    df_merged, excluded = load_merge_frames_from_files(source_files, progress_cb=progress_cb)

    # 2) Save merged (Line 기준 결과)
    push_progress(72, "병합 결과 저장 중(1/2)...")
    merged_file = os.path.join(destination_folder, generate_simple_filename(base_name, 1))
    df_merged.to_excel(merged_file, index=False)
    logger.info("Saved merged file: %s", merged_file)

    # 3) Sort + Save
    sorted_file = None
    dynamic_line = None
    dynamic_sorted = None

    if sort_criteria not in df_merged.columns:
        # sort 컬럼 없으면 종료(기존 동작: 경고 후 종료)
        push_progress(100, f"정렬 기준 열 '{sort_criteria}' 없음")
        return MergeResult(
            merged_file_path=merged_file,
            sorted_file_path=None,
            excluded_sheets=excluded,
            dynamic_chart_line=None,
            dynamic_chart_sorted=None,
        )

    push_progress(82, "정렬 및 저장 중(2/2)...")
    df_sorted = sort_and_reorder_by_column(df_merged, sort_criteria)
    sorted_file = os.path.join(destination_folder, generate_simple_filename(base_name, 2))
    df_sorted.to_excel(sorted_file, index=False)
    logger.info("Saved sorted file: %s", sorted_file)

    # 4) Excel chart
    push_progress(90, "엑셀 차트 추가 중...")
    try:
        add_chart_to_excel(sorted_file)
    except Exception:
        logger.exception("Failed to add chart to Excel (ignored)")

    # 5) Plotly htmls (Line 기반 + Sorted 기반 2개)
    chart_title = f"Chart based on {sort_criteria}"

    push_progress(94, "동적 그래프 생성 중(1/2)...")
    try:
        df_line = pd.read_excel(merged_file)
        x_axis_line = df_line.columns[0]  # 보통 Line
        dynamic_line = plot_from_dataframe(
            df=df_line,
            x_axis_column=x_axis_line,
            title=chart_title,
            destination_folder=destination_folder,
            out_name=f"{base_name}_line_chart.html",
        )
        logger.info("Saved dynamic line chart: %s", dynamic_line)
    except Exception:
        logger.exception("Failed to create line dynamic chart (ignored)")

    push_progress(97, "동적 그래프 생성 중(2/2)...")
    try:
        df_sort = pd.read_excel(sorted_file)
        x_axis_sorted = df_sort.columns[0]  # sort_criteria가 맨 앞으로 오므로 보통 sort_criteria
        dynamic_sorted = plot_from_dataframe(
            df=df_sort,
            x_axis_column=x_axis_sorted,
            title=chart_title,
            destination_folder=destination_folder,
            out_name=f"{base_name}_sorted_chart.html",
        )
        logger.info("Saved dynamic sorted chart: %s", dynamic_sorted)
    except Exception:
        logger.exception("Failed to create sorted dynamic chart (ignored)")

    push_progress(100, "완료!")

    return MergeResult(
        merged_file_path=merged_file,
        sorted_file_path=sorted_file,
        excluded_sheets=excluded,
        dynamic_chart_line=dynamic_line,
        dynamic_chart_sorted=dynamic_sorted,
    )


def start_worker(
    source_files: List[str],
    destination_folder: str,
    sort_criteria: str,
    progress_q: "queue.Queue[ProgressEvent]",
    result_q: "queue.Queue[MergeResult]",
    error_q: "queue.Queue[Exception]",
) -> None:
    """워커 스레드 엔트리."""
    try:
        result = merge_pipeline(source_files, destination_folder, sort_criteria, progress_q)
        result_q.put(result)
    except Exception as e:
        logger.exception("Worker failed")
        error_q.put(e)


# -------------------------
# Main UI
# -------------------------

def main() -> None:
    setup_logging()
    logger.info("Program started.")

    root = tk.Tk()
    root.title("Merge Excel (Refactored)")
    root.geometry("420x140")

    # 진행률 UI
    progress_label = ttk.Label(root, text="진행률: 0%")
    progress_label.pack(pady=10)

    progress_bar = ttk.Progressbar(root, length=360, mode="determinate", maximum=100)
    progress_bar.pack(pady=5)

    # 큐들
    progress_q: "queue.Queue[ProgressEvent]" = queue.Queue()
    result_q: "queue.Queue[MergeResult]" = queue.Queue()
    error_q: "queue.Queue[Exception]" = queue.Queue()

    # 입력 단계(메인 스레드에서)
    root.withdraw()  # 입력 다이얼로그 동안 메인 창 숨김

    source_files = select_files(root, "병합할 엑셀 파일들을 선택하세요 / Select Excel files to merge")
    if not source_files:
        messagebox.showwarning("경고", "파일을 선택하지 않았습니다.", parent=root)
        root.destroy()
        return

    destination_folder = select_folder(root, "대상 폴더를 선택하세요 / Select destination folder")
    if not destination_folder:
        messagebox.showwarning("경고", "폴더를 선택하지 않았습니다.", parent=root)
        root.destroy()
        return

    sort_criteria = select_sort_criteria(root)
    if not sort_criteria:
        messagebox.showwarning("경고", "정렬 기준을 입력하지 않았습니다.", parent=root)
        root.destroy()
        return

    messagebox.showinfo("중복 데이터 처리", "중복된 데이터는 마지막 값으로 처리됩니다.", parent=root)

    root.deiconify()

    # 워커 실행
    worker = threading.Thread(
        target=start_worker,
        args=(source_files, destination_folder, sort_criteria, progress_q, result_q, error_q),
        daemon=True,
    )
    worker.start()

    def poll_queues() -> None:
        # 1) progress 업데이트
        try:
            while True:
                ev = progress_q.get_nowait()
                progress_bar["value"] = ev.percent
                progress_label.config(text=f"진행률: {ev.percent}%  |  {ev.message}")
        except queue.Empty:
            pass

        # 2) error 확인
        try:
            err = error_q.get_nowait()
            messagebox.showerror("오류", f"오류가 발생했습니다:\n{err}", parent=root)
            root.destroy()
            return
        except queue.Empty:
            pass

        # 3) 결과 확인
        try:
            result = result_q.get_nowait()
            # sort 컬럼 없어서 종료되는 경우 처리
            if result.sorted_file_path is None:
                messagebox.showwarning(
                    "경고",
                    f"열 '{sort_criteria}'을(를) 찾을 수 없습니다.\n병합 파일만 저장되었습니다:\n{result.merged_file_path}",
                    parent=root,
                )
            else:
                msg_lines = [
                    "완료!",
                    f"- 병합 파일: {result.merged_file_path}",
                    f"- 정렬 파일: {result.sorted_file_path}",
                ]
                if result.dynamic_chart_line:
                    msg_lines.append(f"- Line 그래프: {result.dynamic_chart_line}")
                if result.dynamic_chart_sorted:
                    msg_lines.append(f"- Sorted 그래프: {result.dynamic_chart_sorted}")
                if result.excluded_sheets:
                    msg_lines.append(f"- 제외된 시트 수: {len(result.excluded_sheets)}")
                messagebox.showinfo("완료", "\n".join(msg_lines), parent=root)

            root.destroy()
            return
        except queue.Empty:
            pass

        # 계속 폴링
        root.after(80, poll_queues)  # type: ignore

    root.after(80, poll_queues)  # type: ignore

    def on_close() -> None:
        # daemon thread라 창 닫으면 종료됨
        root.destroy()
        sys.exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)

    root.mainloop()


if __name__ == "__main__":
    main()
